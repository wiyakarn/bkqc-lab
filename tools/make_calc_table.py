"""Build ตารางคำนวณสูตรอัตโนมัติ.xlsx : product master -> entry log -> auto summary."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

OUT = "/home/user/bkqc-lab/ตารางคำนวณสูตรอัตโนมัติ.xlsx"
N = 200          # master rows
M = 1000         # entry rows
FONT = "Sarabun"
HDR = PatternFill("solid", fgColor="1D6F42"); INP = PatternFill("solid", fgColor="FFF8E1")
AUTO = PatternFill("solid", fgColor="E8F0FA")
thin = Side(style="thin", color="D0CEC4"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def hdr(ws, row, labels, widths):
    for i, (l, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row, i, l); c.fill = HDR; c.border = BOX
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = Font(name=FONT, bold=True, size=15, color="1D6F42")
    ws["A2"] = sub; ws["A2"].font = Font(name=FONT, size=10, color="6B6860")

wb = Workbook()

# ---------- 1) Master ----------
ms = wb.active; ms.title = "สินค้า"
title(ms, "ทะเบียนสินค้า (Master)", "กรอกเฉพาะช่องสีเหลือง — เพิ่มสินค้าใหม่ที่นี่ แล้วชื่อจะขึ้นใน dropdown และตารางสรุปอัตโนมัติ")
hdr(ms, 4, ["SKU", "ชื่อสินค้า", "หมวด", "หน่วย"], [18, 58, 16, 10])
sample = [
 ("SZ-1000032688","VEG-FROZEN BROCCOLI-1KG/BAG-10BAG/CASE","Veg","BAG"),
 ("SZ-1000031958","FRENCH FRIES-XLF 10MM SKIN OFF-2.5 KG/BAG-5BAG/CASE","Potato","CS"),
 ("SZ-1000037051","HASH BROWN-1KG/BAG-10BAG/CASE-CASE","Potato","CS"),
]
for r in range(5, 5 + N):
    for c in range(1, 5):
        cell = ms.cell(r, c); cell.fill = INP; cell.border = BOX; cell.font = Font(name=FONT)
    if r - 5 < len(sample):
        for c, v in enumerate(sample[r - 5], 1): ms.cell(r, c, v)
ms.freeze_panes = "A5"
ML = f"สินค้า!$B$5:$B${4+N}"

# ---------- 2) Entry ----------
en = wb.create_sheet("บันทึก")
title(en, "บันทึกการส่งตรวจ", "เลือกชื่อสินค้าจาก dropdown (ช่องเหลือง) → ช่องฟ้าคำนวณเองทั้งหมด")
labels = ["วันที่ส่ง","PO","ชื่อสินค้า","SKU\n(auto)","หมวด\n(auto)","หน่วย\n(auto)",
          "จำนวนที่มาทั้งหมด","จำนวนที่ส่งตรวจ","% ที่สุ่มตรวจ\n(auto)","สถานะ QA","COA",
          "เดือน\n(auto)","วันที่ค้าง\n(auto)"]
hdr(en, 4, labels, [12,16,52,16,10,8,14,14,12,12,8,10,10])
dv_prod = DataValidation(type="list", formula1=f"={ML}", allow_blank=True,
                         showErrorMessage=True, errorTitle="ไม่พบสินค้า",
                         error="เลือกจากรายการ หรือเพิ่มสินค้าในชีต 'สินค้า' ก่อน")
dv_stat = DataValidation(type="list", formula1='"Done,Pending"', allow_blank=True)
dv_coa = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, error="กรอกตัวเลขเท่านั้น")
for d in (dv_prod, dv_stat, dv_coa, dv_num): en.add_data_validation(d)
last = 4 + M
dv_prod.add(f"C5:C{last}"); dv_stat.add(f"J5:J{last}"); dv_coa.add(f"K5:K{last}")
dv_num.add(f"G5:H{last}")
look = lambda col, r: (f'=IF($C{r}="","",IFERROR(INDEX(สินค้า!${col}$5:${col}${4+N},'
                       f'MATCH($C{r},สินค้า!$B$5:$B${4+N},0)),"ไม่พบสินค้า"))')
demo = [(dt.date(2026,1,6),"C102-00003267",sample[0][1],7700,10,"Done","Yes"),
        (dt.date(2026,1,7),"F106-00073060",sample[1][1],1820,3,"Done","Yes"),
        (dt.date(2026,2,3),"F106-00074001",sample[1][1],900,1,"Pending","No"),
        (dt.date(2026,2,10),"H101-00005511",sample[2][1],2000,2,"Pending","No")]
for r in range(5, last + 1):
    for c in range(1, 14):
        cell = en.cell(r, c); cell.border = BOX; cell.font = Font(name=FONT)
        cell.fill = AUTO if c in (4,5,6,9,12,13) else INP
    en.cell(r, 1).number_format = "dd/mm/yyyy"
    en.cell(r, 4, look("A", r)); en.cell(r, 5, look("C", r)); en.cell(r, 6, look("D", r))
    en.cell(r, 9, f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(H{r}),G{r}>0),H{r}/G{r},"")'); en.cell(r, 9).number_format = "0.00%"
    en.cell(r, 12, f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))')
    en.cell(r, 13, f'=IF(AND(A{r}<>"",J{r}="Pending"),TODAY()-A{r},"")')
    if r - 5 < len(demo):
        for c, v in zip((1,2,3,7,8,10,11), demo[r - 5]): en.cell(r, c, v)
en.freeze_panes = "D5"; en.auto_filter.ref = f"A4:M{last}"
en.conditional_formatting.add(f"J5:J{last}", CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", bgColor="C6EFCE")))
en.conditional_formatting.add(f"J5:J{last}", CellIsRule(operator="equal", formula=['"Pending"'], fill=PatternFill("solid", bgColor="FFEB9C")))
en.conditional_formatting.add(f"M5:M{last}", CellIsRule(operator="greaterThan", formula=["14"], font=Font(color="C0392B", bold=True)))

# ---------- 3) Summary ----------
sm = wb.create_sheet("สรุป", 0)
title(sm, "สรุปผลรวมตามสินค้า (อัตโนมัติ)", "ไม่ต้องกรอก — ดึงจากชีต 'สินค้า' และ 'บันทึก' ทุกครั้งที่มีการแก้ไข")
E = lambda col: f"บันทึก!${col}$5:${col}${last}"
kpi = [("จำนวนครั้งที่ส่งตรวจ", f'=COUNTA({E("C")})'),
       ("Done", f'=COUNTIF({E("J")},"Done")'),
       ("Pending", f'=COUNTIF({E("J")},"Pending")'),
       ("% Done", '=IF(B4=0,0,C4/B4)'),
       ("COA ครบ (Yes)", f'=COUNTIF({E("K")},"Yes")')]
for i, (l, f) in enumerate(kpi):
    c = sm.cell(4, 2 + i); c.value = f.replace("B4", "B4")  # placeholders fixed below
# KPI layout: labels row 3, values row 4 in A..E
for i, (l, f) in enumerate(kpi):
    lc = sm.cell(3, 1 + i, l); lc.font = Font(name=FONT, size=9, color="6B6860"); lc.alignment = Alignment(horizontal="center")
    vc = sm.cell(4, 1 + i); vc.font = Font(name=FONT, bold=True, size=16, color="1D6F42"); vc.alignment = Alignment(horizontal="center")
sm["A4"] = kpi[0][1]; sm["B4"] = kpi[1][1]; sm["C4"] = kpi[2][1]
sm["D4"] = '=IF(A4=0,0,B4/A4)'; sm["D4"].number_format = "0%"; sm["E4"] = kpi[4][1]
for col in "ABCDE":
    sm[f"{col}3"].fill = PatternFill("solid", fgColor="F4F3EF"); sm[f"{col}4"].fill = PatternFill("solid", fgColor="F4F3EF")
for r in range(0, 0): pass
H = 6
hdr(sm, H, ["ชื่อสินค้า","SKU","หมวด","จำนวนครั้ง\nที่ส่งตรวจ","รวมที่มา","รวมที่ส่งตรวจ","หน่วย","Done","Pending","COA (Yes)","% Done","ส่งล่าสุด"],
    [56,16,10,12,12,12,8,8,10,10,9,12])
sm.column_dimensions["A"].width = 56
for i in range(N):
    r = H + 1 + i; m = 5 + i
    sm.cell(r, 1, f'=IF(สินค้า!B{m}="","",สินค้า!B{m})')
    sm.cell(r, 2, f'=IF($A{r}="","",สินค้า!A{m})')
    sm.cell(r, 3, f'=IF($A{r}="","",สินค้า!C{m})')
    sm.cell(r, 4, f'=IF($A{r}="","",COUNTIF({E("C")},$A{r}))')
    sm.cell(r, 5, f'=IF($A{r}="","",SUMIFS({E("G")},{E("C")},$A{r}))')
    sm.cell(r, 6, f'=IF($A{r}="","",SUMIFS({E("H")},{E("C")},$A{r}))')
    sm.cell(r, 7, f'=IF($A{r}="","",สินค้า!D{m})')
    sm.cell(r, 8, f'=IF($A{r}="","",COUNTIFS({E("C")},$A{r},{E("J")},"Done"))')
    sm.cell(r, 9, f'=IF($A{r}="","",COUNTIFS({E("C")},$A{r},{E("J")},"Pending"))')
    sm.cell(r,10, f'=IF($A{r}="","",COUNTIFS({E("C")},$A{r},{E("K")},"Yes"))')
    sm.cell(r,11, f'=IF(OR($A{r}="",D{r}=0),"",H{r}/D{r})'); sm.cell(r,11).number_format = "0%"
    sm.cell(r,12, f'=IF(OR($A{r}="",D{r}=0),"",MAXIFS({E("A")},{E("C")},$A{r}))'); sm.cell(r,12).number_format = "dd/mm/yyyy"
    for c in range(1, 13):
        sm.cell(r, c).font = Font(name=FONT); sm.cell(r, c).border = BOX
        if c in (5, 6): sm.cell(r, c).number_format = "#,##0"
tr = H + 1 + N
sm.cell(tr, 1, "รวมทั้งหมด")
for c, col in ((4,"D"),(5,"E"),(6,"F"),(8,"H"),(9,"I"),(10,"J")):
    sm.cell(tr, c, f"=SUM({col}{H+1}:{col}{tr-1})").number_format = "#,##0"
for c in range(1, 13):
    cell = sm.cell(tr, c); cell.font = Font(name=FONT, bold=True); cell.fill = PatternFill("solid", fgColor="E8F3ED"); cell.border = BOX
sm.conditional_formatting.add(f"I{H+1}:I{tr-1}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="9A6200", bold=True)))
sm.freeze_panes = f"B{H+1}"

# ---------- 4) Monthly ----------
mo = wb.create_sheet("สรุปรายเดือน", 1)
title(mo, "สรุปรายเดือน (อัตโนมัติ)", "เปลี่ยนปีที่ช่อง B3 ได้")
mo["A3"] = "ปี (ค.ศ.)"; mo["B3"] = 2026; mo["B3"].fill = INP
hdr(mo, 5, ["เดือน","จำนวนครั้ง","Done","Pending","รวมที่ส่งตรวจ"], [14,14,10,10,16])
for i in range(12):
    r = 6 + i
    mo.cell(r, 1, f'=TEXT(DATE($B$3,{i+1},1),"yyyy-mm")')
    mo.cell(r, 2, f'=COUNTIF({E("L")},$A{r})')
    mo.cell(r, 3, f'=COUNTIFS({E("L")},$A{r},{E("J")},"Done")')
    mo.cell(r, 4, f'=COUNTIFS({E("L")},$A{r},{E("J")},"Pending")')
    mo.cell(r, 5, f'=SUMIFS({E("H")},{E("L")},$A{r})')
    for c in range(1, 6): mo.cell(r, c).border = BOX; mo.cell(r, c).font = Font(name=FONT)
mo.cell(18, 1, "รวม")
for c, col in ((2,"B"),(3,"C"),(4,"D"),(5,"E")): mo.cell(18, c, f"=SUM({col}6:{col}17)")
for c in range(1, 6): mo.cell(18, c).font = Font(name=FONT, bold=True); mo.cell(18, c).fill = PatternFill("solid", fgColor="E8F3ED")

# ---------- 5) Guide ----------
gd = wb.create_sheet("วิธีใช้")
gd.column_dimensions["A"].width = 100
for i, t in enumerate([
 "วิธีใช้งาน",
 "1) ชีต 'สินค้า' – ใส่ SKU / ชื่อสินค้า / หมวด / หน่วย (ช่องเหลือง) รองรับ 200 รายการ",
 "2) ชีต 'บันทึก' – กรอกวันที่ PO เลือกชื่อสินค้าจาก dropdown ใส่จำนวน และสถานะ  (SKU หมวด หน่วย % เดือน วันค้าง คำนวณเอง)",
 "3) ชีต 'สรุป' และ 'สรุปรายเดือน' – รวมข้อมูลตามสินค้า/เดือนอัตโนมัติ ห้ามพิมพ์ทับ",
 "สูตรหลักที่ใช้: INDEX/MATCH (ดึง SKU), COUNTIFS/SUMIFS (รวม), MAXIFS (วันที่ล่าสุด), TODAY (วันที่ค้าง)",
 "ช่องสีเหลือง = กรอกเอง | ช่องสีฟ้า = สูตรอัตโนมัติ",
 "ตัวอย่างข้อมูลเป็นข้อมูลสมมติ ลบแถวตัวอย่างในชีต 'บันทึก' ก่อนใช้งานจริงได้"], 1):
    gd.cell(i, 1, t).font = Font(name=FONT, bold=(i == 1), size=13 if i == 1 else 11)
wb.save(OUT); print("saved", OUT)
